"""The workbook itself needs Excel to calculate, so these tests check what can be checked
without it: that the build succeeds, the structure is right, the formulas point at the right
ranges, and bad input is refused with a useful message."""

import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "build_report.py"

INV = "invoice_id,account_code,po_number,invoice_amount,invoice_date\n"
PAY = "payment_id,account_code,po_number,payment_amount,payment_date,payment_ref\n"


def build(tmp_path, invoices=None, payments=None):
    args = [sys.executable, str(SCRIPT), "--out", str(tmp_path / "out.xlsx")]
    if invoices is not None:
        (tmp_path / "inv.csv").write_text(invoices)
        args += ["--invoices", str(tmp_path / "inv.csv")]
    if payments is not None:
        (tmp_path / "pay.csv").write_text(payments)
        args += ["--payments", str(tmp_path / "pay.csv")]
    return subprocess.run(args, capture_output=True, text=True)


def test_default_sample_builds_the_expected_sheets(tmp_path):
    r = build(tmp_path)
    assert r.returncode == 0, r.stderr
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["Summary", "Reconciliation", "Payments"]
    assert wb["Reconciliation"].max_row == 101          # 100 invoices + header
    assert wb["Payments"].max_row == 106                # 105 payments + header
    assert "SAMPLE DATA" in wb["Summary"]["A2"].value


def test_formulas_point_at_the_right_ranges(tmp_path):
    build(tmp_path)
    ws = load_workbook(tmp_path / "out.xlsx")["Reconciliation"]
    assert ws["F2"].value.startswith("=SUMIFS(Payments!$E$2:$E$106")
    assert ws["G2"].value.startswith("=COUNTIFS(Payments!$B$2:$B$106")
    assert ws["I2"].value.startswith("=IF(G2=0,\"Unpaid\"")
    assert "Summary!$B$4" in ws["I2"].value                # tolerance comes from the input cell


def test_your_own_files_are_used_and_labelled(tmp_path):
    inv = INV + "1,AC1,PO1,100.00,2026-06-01\n2,AC1,PO2,50.00,2026-06-02\n"
    pay = PAY + "1,AC1,PO1,100.00,2026-06-10,R1\n"
    assert build(tmp_path, inv, pay).returncode == 0
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb["Reconciliation"].max_row == 3 and wb["Payments"].max_row == 2
    assert "inv.csv" in wb["Summary"]["A2"].value and "SAMPLE" not in wb["Summary"]["A2"].value


def test_missing_columns_are_named(tmp_path):
    r = build(tmp_path, invoices="a,b\n1,2\n", payments=PAY + "1,AC1,PO1,1,2026-01-01,R\n")
    assert r.returncode != 0
    assert "missing column(s): invoice_id" in r.stderr and "Columns found: a, b" in r.stderr


def test_non_numeric_amount_is_refused_with_its_row(tmp_path):
    inv = INV + "1,AC1,PO1,100,2026-06-01\n2,AC1,PO2,TBD,2026-06-02\n"
    r = build(tmp_path, inv, PAY + "1,AC1,PO1,100,2026-06-10,R\n")
    assert r.returncode != 0 and "non-numeric" in r.stderr and "data row 2" in r.stderr


def test_missing_file_is_reported(tmp_path):
    r = subprocess.run([sys.executable, str(SCRIPT), "--invoices", str(tmp_path / "nope.csv"),
                        "--out", str(tmp_path / "o.xlsx")], capture_output=True, text=True)
    assert r.returncode != 0 and "not found" in r.stderr


def test_excel_bom_csv_is_accepted(tmp_path):
    (tmp_path / "inv.csv").write_bytes(("﻿" + INV + "1,AC1,PO1,10,2026-06-01\n").encode("utf-8"))
    (tmp_path / "pay.csv").write_text(PAY + "1,AC1,PO1,10,2026-06-02,R\n")
    r = subprocess.run([sys.executable, str(SCRIPT), "--invoices", str(tmp_path / "inv.csv"),
                        "--payments", str(tmp_path / "pay.csv"), "--out", str(tmp_path / "o.xlsx")],
                       capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
