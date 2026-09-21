import argparse
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

INVOICE_COLUMNS = ["invoice_id", "account_code", "po_number", "invoice_amount", "invoice_date"]
PAYMENT_COLUMNS = ["payment_id", "account_code", "po_number", "payment_amount", "payment_date", "payment_ref"]


def parse_args():
    ap = argparse.ArgumentParser(
        description="Build the payout reconciliation workbook from an invoice CSV and a payment CSV.")
    ap.add_argument("--invoices", default=os.path.join(HERE, "data", "invoices.csv"),
                    help="invoice CSV with columns: " + ", ".join(INVOICE_COLUMNS))
    ap.add_argument("--payments", default=os.path.join(HERE, "data", "payments.csv"),
                    help="payment CSV with columns: " + ", ".join(PAYMENT_COLUMNS))
    ap.add_argument("--out", default=os.path.join(HERE, "Sample_Payout_Reconciliation_Report.xlsx"),
                    help="Excel file to write")
    return ap.parse_args()


def read_csv(path, required, kind):
    if not os.path.exists(path):
        sys.exit(f"Error: {kind} file not found: {path}")
    df = pd.read_csv(path, dtype={"po_number": str}, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in required if c not in df.columns]
    if missing:
        sys.exit(f"Error: {os.path.basename(path)} is missing column(s): {', '.join(missing)}. "
                 f"Columns found: {', '.join(df.columns)}. "
                 f"Expected: {', '.join(required)}.")
    df["po_number"] = df["po_number"].astype(str).str.strip()
    for col in (required[3],):                              # the amount column must be numeric
        bad = pd.to_numeric(df[col], errors="coerce").isna()
        if bad.any():
            sys.exit(f"Error: {os.path.basename(path)} has {int(bad.sum())} non-numeric value(s) in "
                     f"'{col}' (first at data row {int(bad.idxmax()) + 1}).")
    return df


ARGS = parse_args()
OUT = ARGS.out
os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
USING_SAMPLE = (os.path.abspath(ARGS.invoices) == os.path.join(HERE, "data", "invoices.csv")
                and os.path.abspath(ARGS.payments) == os.path.join(HERE, "data", "payments.csv"))

invoices = read_csv(ARGS.invoices, INVOICE_COLUMNS, "invoice")
payments = read_csv(ARGS.payments, PAYMENT_COLUMNS, "payment")
invoices["invoice_date"] = pd.to_datetime(invoices["invoice_date"]).dt.date
payments["payment_date"] = pd.to_datetime(payments["payment_date"]).dt.date

n_inv, n_pay = len(invoices), len(payments)
inv_last, pay_last = n_inv + 1, n_pay + 1

FONT = "Arial"
NAVY = "1F3864"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_note = Font(name=FONT, size=9, italic=True, color="595959")
f_input = Font(name=FONT, size=10, color="0000FF")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_total = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '$#,##0.00;($#,##0.00);-'
MONEY0 = '$#,##0;($#,##0);-'

wb = Workbook()

# ---------------------------------------------------------------- Reconciliation
rec = wb.active
rec.title = "Reconciliation"
heads = ["Invoice ID", "Account", "PO Number", "Invoice Date", "Invoice Amount",
         "Total Paid", "Payments", "Difference", "Status"]
rec.append(heads)
for i, r in enumerate(invoices.itertuples(index=False), start=2):
    rec.append([r.invoice_id, r.account_code, r.po_number, r.invoice_date, r.invoice_amount])
    rec[f"F{i}"] = (f"=SUMIFS(Payments!$E$2:$E${pay_last},Payments!$B$2:$B${pay_last},B{i},"
                    f"Payments!$C$2:$C${pay_last},C{i})")
    rec[f"G{i}"] = (f"=COUNTIFS(Payments!$B$2:$B${pay_last},B{i},"
                    f"Payments!$C$2:$C${pay_last},C{i})")
    rec[f"H{i}"] = f"=ROUND(E{i}-F{i},2)"
    tol = "Summary!$B$4"
    rec[f"I{i}"] = (f'=IF(G{i}=0,"Unpaid",IF(ABS(H{i})<={tol},"Paid",IF(H{i}>{tol},"Short Pay",'
                    f'IF(AND(G{i}>=2,OR(ABS(F{i}-2*E{i})<={tol},ABS(F{i}-3*E{i})<={tol},'
                    f'ABS(F{i}-4*E{i})<={tol})),"Duplicate","Overpaid"))))')

for c in range(1, len(heads) + 1):
    cell = rec.cell(row=1, column=c)
    cell.font, cell.fill, cell.border = f_head, fill_head, box
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in rec.iter_rows(min_row=2, max_row=inv_last, max_col=len(heads)):
    for cell in row:
        cell.font, cell.border = f_base, box
    row[3].number_format = "yyyy-mm-dd"
    for idx in (4, 5, 7):
        row[idx].number_format = MONEY
    row[6].alignment = Alignment(horizontal="center")
    row[8].alignment = Alignment(horizontal="center")
for col, w in zip("ABCDEFGHI", [13, 11, 14, 14, 16, 14, 12, 14, 12]):
    rec.column_dimensions[col].width = w
rec.freeze_panes = "A2"
rec.auto_filter.ref = f"A1:I{inv_last}"
rng = f"I2:I{inv_last}"
for label, color in [("Unpaid", "F8CBAD"), ("Short Pay", "FFE699"), ("Overpaid", "BDD7EE"),
                     ("Duplicate", "D9B3FF"), ("Paid", "C6E0B4")]:
    rec.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=[f'"{label}"'],
                        fill=PatternFill("solid", bgColor=color, fgColor=color)))

# ---------------------------------------------------------------- Payments
pay = wb.create_sheet("Payments")
pheads = ["Payment ID", "Account", "PO Number", "Payment Date", "Payment Amount",
          "Reference", "Invoice Match"]
pay.append(pheads)
for i, r in enumerate(payments.itertuples(index=False), start=2):
    pay.append([r.payment_id, r.account_code, r.po_number, r.payment_date,
                r.payment_amount, r.payment_ref])
    pay[f"G{i}"] = (f'=IF(COUNTIFS(Reconciliation!$B$2:$B${inv_last},B{i},'
                    f'Reconciliation!$C$2:$C${inv_last},C{i})=0,"Unapplied","Applied")')
for c in range(1, len(pheads) + 1):
    cell = pay.cell(row=1, column=c)
    cell.font, cell.fill, cell.border = f_head, fill_head, box
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in pay.iter_rows(min_row=2, max_row=pay_last, max_col=len(pheads)):
    for cell in row:
        cell.font, cell.border = f_base, box
    row[3].number_format = "yyyy-mm-dd"
    row[4].number_format = MONEY
    row[6].alignment = Alignment(horizontal="center")
for col, w in zip("ABCDEFG", [11, 10, 13, 13, 15, 12, 14]):
    pay.column_dimensions[col].width = w
pay.freeze_panes = "A2"
pay.auto_filter.ref = f"A1:G{pay_last}"
pay.conditional_formatting.add(
    f"G2:G{pay_last}",
    CellIsRule(operator="equal", formula=['"Unapplied"'],
               fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))

# ---------------------------------------------------------------- Summary
s = wb.create_sheet("Summary", 0)
s.sheet_view.showGridLines = False
s["A1"] = "Marketplace Payout Reconciliation Report"
s["A1"].font = f_title
s["A2"] = (
    "SAMPLE DATA: synthetic invoices and payments generated for demonstration. "
    "Not real client data." if USING_SAMPLE else
    f"Built from {os.path.basename(ARGS.invoices)} and {os.path.basename(ARGS.payments)}.")
s["A2"].font = f_note
s["A4"] = "Match tolerance ($)"
s["A4"].font = f_bold
s["B4"] = 0.01
s["B4"].font, s["B4"].fill, s["B4"].number_format = f_input, fill_input, "$0.00"
s["C4"] = "Input: a difference within this amount counts as Paid."
s["C4"].font = f_note

hdr = ["Status", "Invoices", "% of Invoices", "Invoiced ($)", "Paid ($)", "Difference ($)"]
for c, h in enumerate(hdr, start=1):
    cell = s.cell(row=6, column=c, value=h)
    cell.font, cell.fill, cell.border = f_head, fill_head, box
    cell.alignment = Alignment(horizontal="center", vertical="center")
statuses = ["Paid", "Unpaid", "Short Pay", "Overpaid", "Duplicate"]
R = f"Reconciliation!$I$2:$I${inv_last}"
for i, st in enumerate(statuses, start=7):
    s[f"A{i}"] = st
    s[f"B{i}"] = f"=COUNTIF({R},A{i})"
    s[f"C{i}"] = f"=IF($B$12=0,0,B{i}/$B$12)"
    s[f"D{i}"] = f"=SUMIF({R},A{i},Reconciliation!$E$2:$E${inv_last})"
    s[f"E{i}"] = f"=SUMIF({R},A{i},Reconciliation!$F$2:$F${inv_last})"
    s[f"F{i}"] = f"=SUMIF({R},A{i},Reconciliation!$H$2:$H${inv_last})"
s["A12"] = "Total"
for col in "BDEF":
    s[f"{col}12"] = f"=SUM({col}7:{col}11)"
s["C12"] = "=SUM(C7:C11)"
for r in range(7, 13):
    for c in range(1, 7):
        cell = s.cell(row=r, column=c)
        cell.font = f_bold if r == 12 else f_base
        cell.border = box
        if r == 12:
            cell.fill = fill_total
    s[f"B{r}"].number_format = "#,##0"
    s[f"C{r}"].number_format = "0.0%"
    for col in "DEF":
        s[f"{col}{r}"].number_format = MONEY0

s["A14"] = "Key findings"
s["A14"].font = Font(name=FONT, size=11, bold=True, color=NAVY)
findings = [
    ("Money owed (Unpaid + Short Pay) ($)", "=F8+F9", MONEY0),
    ("Excess received (Overpaid + Duplicate) ($)", "=-(F10+F11)", MONEY0),
    ("Unapplied payments (count)", f'=COUNTIF(Payments!$G$2:$G${pay_last},"Unapplied")', "#,##0"),
    ("Unapplied payments ($)",
     f'=SUMIF(Payments!$G$2:$G${pay_last},"Unapplied",Payments!$E$2:$E${pay_last})', MONEY0),
]
for i, (label, formula, fmt) in enumerate(findings, start=15):
    s[f"A{i}"] = label
    s[f"B{i}"] = formula
    s[f"A{i}"].font = f_base
    s[f"B{i}"].font = f_bold
    s[f"B{i}"].number_format = fmt
    s[f"A{i}"].border = box
    s[f"B{i}"].border = box

s["A20"] = ("How to read this: the Reconciliation sheet matches every invoice to its payments by "
            "Account + PO Number. Filter the Status column to see exactly which invoices to chase.")
s["A20"].font = f_note
s["A21"] = ("Status rules: Unpaid = no payment; Paid = matches within tolerance; Short Pay = paid less; "
            "Overpaid = paid more; Duplicate = 2+ payments totalling 2x, 3x or 4x the invoice.")
s["A21"].font = f_note

for col, w in zip("ABCDEF", [42, 12, 14, 16, 16, 16]):
    s.column_dimensions[col].width = w

from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint

chart = BarChart()
chart.type = "col"
chart.title = "Invoices by reconciliation status"
chart.legend = None
chart.varyColors = False
chart.add_data(Reference(s, min_col=2, min_row=6, max_row=11), titles_from_data=True)
chart.set_categories(Reference(s, min_col=1, min_row=7, max_row=11))
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.y_axis.majorGridlines = None
ser = chart.series[0]
ser.dLbls = DataLabelList()
ser.dLbls.showVal = True
ser.dLbls.showSerName = ser.dLbls.showCatName = ser.dLbls.showLegendKey = False
for idx, color in enumerate(["70AD47", "F4B183", "FFC000", "5B9BD5", "9E6BD1"]):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties.solidFill = color
    pt.graphicalProperties.line.solidFill = color
    ser.dPt.append(pt)
chart.gapWidth = 60
chart.height, chart.width = 7.5, 13
s.add_chart(chart, "H4")

wb.save(OUT)
print("saved", OUT, "| invoices:", n_inv, "| payments:", n_pay)

